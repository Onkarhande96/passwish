import os
import re
import json
import random
import subprocess
import urllib.parse
import datetime
import time
import threading
import openpyxl
from dotenv import load_dotenv
load_dotenv()
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from functools import wraps
from flask import (
    Flask, render_template, request, jsonify, send_file, send_from_directory,
    session as flask_session, redirect, url_for, has_request_context
)
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash
from sqlalchemy import or_, and_, func

from database import (
    init_db, SessionLocal, Customer, MessageTemplate, SentLog, Setting,
    User, BackupLog, DB_URL, is_sqlite
)
from importer import process_excel_file, get_preview_data, clean_phone
from backup_manager import create_backup, restore_backup, list_backups, BACKUP_DIR
import license_manager
from license_manager import (
    get_machine_id, verify_license_token, save_license,
    is_software_activated, get_license_details
)
from reminder_manager import start_reminder_daemon, trigger_desktop_notification

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'passwish-crm-secure-key-2026')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50 MB max file size
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Initialize DB tables & seed templates on boot
init_db()

# Start background morning reminder daemon
start_reminder_daemon(SessionLocal, Customer, Setting)

# --- Hardware-Locked License Enforcement ---
@app.before_request
def enforce_machine_license():
    # Allow static files and activation endpoints without activation
    allowed_prefixes = (
        '/static',
        '/uploads',
        '/activate',
        '/api/license',
        '/health'
    )
    for p in allowed_prefixes:
        if request.path.startswith(p):
            return None

    activated, msg, _ = is_software_activated()
    if not activated:
        if request.path.startswith('/api/'):
            return jsonify({
                'error': 'Software not activated. Valid license key required for this machine.',
                'machine_id': get_machine_id(),
                'status': 'unlicensed'
            }), 403
        return redirect('/activate')

# --- Auth Decorators ---
def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        uid = flask_session.get('user_id')
        if not uid:
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Unauthorized. Please login.'}), 401
            return redirect(url_for('login_page'))

        # Check license validity for non-admin accounts
        if flask_session.get('role') != 'admin':
            db = SessionLocal()
            try:
                user = db.query(User).filter_by(id=uid).first()
                if not user or not user.is_active or user.get_license_info().get('is_expired'):
                    flask_session.clear()
                    if request.path.startswith('/api/'):
                        return jsonify({'error': 'Your license has expired or your account has been deactivated. Please contact admin.'}), 403
                    return redirect(url_for('login_page', expired=1))
            finally:
                db.close()

        return f(*args, **kwargs)
    return decorated_function

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not flask_session.get('user_id'):
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Unauthorized. Please login.'}), 401
            return redirect(url_for('admin_login_page'))
        if flask_session.get('role') != 'admin':
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Forbidden: Administrator privileges required.'}), 403
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function

# --- Tenant Isolation & Scoped Helper Functions ---
def get_user_setting(user_id, key, default=''):
    db = SessionLocal()
    try:
        if user_id:
            s = db.query(Setting).filter_by(key=f"user_{user_id}_{key}").first()
            if s and s.value:
                return s.value
            # Fallback to agency_name / phone / full_name from User model
            u = db.query(User).filter_by(id=user_id).first()
            if u:
                if key == 'shop_name' and u.agency_name:
                    return u.agency_name
                elif key == 'phone' and u.phone:
                    return u.phone
                elif key == 'owner_name' and u.full_name:
                    return u.full_name

            # DO NOT fallback to global settings for tenant-specific profile fields!
            if key in ('shop_name', 'owner_name', 'phone', 'address'):
                return default

        # For global settings (whatsapp_mode, reminder_enabled, etc.)
        if key in ('shop_name', 'owner_name', 'phone', 'address'):
            return default
        s = db.query(Setting).filter_by(key=key).first()
        return s.value if s else default
    finally:
        db.close()

def get_setting(key, default=''):
    uid = flask_session.get('user_id') if has_request_context() else None
    return get_user_setting(uid, key, default)

def get_tenant_query(session, model=Customer):
    """
    Returns a query strictly scoped to the logged-in tenant/user.
    Non-admins (passport shops) can ONLY see their own records.
    Super Admins can view all or filter by ?client_id=<id>.
    """
    uid = flask_session.get('user_id') if has_request_context() else None
    role = flask_session.get('role', 'staff') if has_request_context() else 'staff'

    query = session.query(model)
    if role != 'admin':
        if hasattr(model, 'created_by_user_id'):
            query = query.filter(model.created_by_user_id == uid)
        elif hasattr(model, 'sent_by_user_id'):
            query = query.filter(model.sent_by_user_id == uid)
    else:
        if has_request_context():
            client_id = request.args.get('client_id')
            if client_id:
                try:
                    cid_int = int(client_id)
                    if hasattr(model, 'created_by_user_id'):
                        query = query.filter(model.created_by_user_id == cid_int)
                    elif hasattr(model, 'sent_by_user_id'):
                        query = query.filter(model.sent_by_user_id == cid_int)
                except Exception:
                    pass
    return query

def validate_mobile_for_queue(phone_str):
    """
    Validates if a phone string is a valid 10-digit Indian mobile number suitable for WhatsApp.
    Returns (is_valid, cleaned_international_phone, reason)
    """
    if not phone_str:
        return False, '', 'Missing phone number'
    digits = re.sub(r'\D', '', str(phone_str))
    if len(digits) == 12 and digits.startswith('91'):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith('0'):
        digits = digits[1:]

    if len(digits) != 10:
        return False, str(phone_str), f"Invalid length ({len(digits)} digits, must be 10 digits)"
    if digits[0] not in ('6', '7', '8', '9'):
        return False, str(phone_str), f"Landline or non-mobile series (starts with '{digits[0]}')"
    if digits == digits[0] * 10:
        return False, str(phone_str), "Invalid repeated digits"

    return True, f"91{digits}", "Valid mobile"

def build_whatsapp_url(phone, message_text):
    is_valid, clean_p, _ = validate_mobile_for_queue(phone)
    if not is_valid:
        clean_p = str(phone).strip().replace('+', '').replace(' ', '').replace('-', '')
    encoded_text = urllib.parse.quote(message_text)
    return f"https://web.whatsapp.com/send?phone={clean_p}&text={encoded_text}"

def render_template_text(template_body, customer, festival_name='', user_id=None):
    uid = user_id or (flask_session.get('user_id') if has_request_context() else None)
    if not uid and customer and hasattr(customer, 'created_by_user_id') and customer.created_by_user_id:
        uid = customer.created_by_user_id

    shop_name = get_user_setting(uid, 'shop_name', '').strip()
    shop_phone = get_user_setting(uid, 'phone', '').strip()
    shop_address = get_user_setting(uid, 'address', '').strip()

    text = template_body
    text = text.replace('{name}', customer.name or '')
    text = text.replace('{phone}', customer.phone or '')
    text = text.replace('{passport_no}', customer.passport_no or '')
    text = text.replace('{expiry_date}', customer.expiry_date or '')
    text = text.replace('{city}', customer.city or '')
    text = text.replace('{shop_name}', shop_name)
    text = text.replace('{shop_phone}', shop_phone)
    text = text.replace('{shop_address}', shop_address)
    text = text.replace('{festival_name}', festival_name)
    return text

# --- Authentication Routes ---
@app.route('/login', methods=['GET', 'POST'])
def login_page():
    if flask_session.get('user_id'):
        if flask_session.get('role') == 'admin':
            return redirect(url_for('admin_portal'))
        return redirect(url_for('index'))

    error = None
    if request.args.get('expired'):
        error = 'तुमचा लायसन्स कालावधी संपला आहे. कृपया ॲडमिनशी संपर्क साधा.'

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()

        db = SessionLocal()
        try:
            user = db.query(User).filter_by(username=username).first()
            if user and user.check_password(password):
                if not getattr(user, 'is_active', True):
                    error = 'This account has been deactivated. Please contact your administrator.'
                else:
                    # Check license validity for non-admin accounts
                    lic_info = user.get_license_info()
                    if user.role != 'admin' and lic_info.get('is_expired'):
                        error = f"तुमचा लायसन्स कालावधी संपला आहे (Expired: {user.valid_until}). कृपया नूतनीकरणासाठी ॲडमिनशी संपर्क साधा."
                    else:
                        flask_session.clear()
                        flask_session['user_id'] = user.id
                        flask_session['username'] = user.username
                        flask_session['role'] = getattr(user, 'role', 'admin' if user.username == 'admin' else 'staff')
                        flask_session['full_name'] = getattr(user, 'full_name', '') or user.username
                        flask_session['agency_name'] = getattr(user, 'agency_name', '') or ''
                        flask_session.permanent = True
                        if flask_session['role'] == 'admin':
                            return redirect(url_for('admin_portal'))
                        return redirect(url_for('index'))
            else:
                error = 'Invalid username or password.'
        finally:
            db.close()

    return render_template('login.html', error=error)

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login_page():
    if flask_session.get('user_id'):
        if flask_session.get('role') == 'admin':
            return redirect(url_for('admin_portal'))
        return redirect(url_for('index'))

    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()

        db = SessionLocal()
        try:
            user = db.query(User).filter_by(username=username).first()
            if user and user.check_password(password):
                if user.role != 'admin':
                    error = 'प्रवेश नाकारला! हे पोर्टल फक्त सुपर ॲडमिनसाठी राखीव आहे. कृपया कन्सल्टंट लॉगिन वापरा.'
                elif not getattr(user, 'is_active', True):
                    error = 'हे ॲडमिन खाते निष्क्रिय (Deactivated) केले आहे.'
                else:
                    flask_session.clear()
                    flask_session['user_id'] = user.id
                    flask_session['username'] = user.username
                    flask_session['role'] = 'admin'
                    flask_session['full_name'] = getattr(user, 'full_name', '') or user.username
                    flask_session['agency_name'] = getattr(user, 'agency_name', '') or ''
                    flask_session.permanent = True
                    return redirect(url_for('admin_portal'))
            else:
                error = 'चुकीचा ॲडमिन आयडी किंवा पासवर्ड (Invalid credentials).'
        finally:
            db.close()

    return render_template('admin_login.html', error=error)

@app.route('/admin')
@admin_required
def admin_portal():
    user_id = flask_session.get('user_id')
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(id=user_id).first()
        return render_template('admin_portal.html', current_user=user.to_dict() if user else {'username': 'admin'})
    finally:
        db.close()

@app.route('/logout')
def logout():
    flask_session.clear()
    return redirect(url_for('login_page'))

@app.route('/api/me', methods=['GET'])
@login_required
def get_current_user_info():
    user_id = flask_session.get('user_id')
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(id=user_id).first()
        if not user:
            return jsonify({'error': 'User not found'}), 404
        return jsonify(user.to_dict())
    finally:
        db.close()

# --- Admin Portal & Licensing APIs ---
@app.route('/api/admin/stats', methods=['GET'])
@admin_required
def get_admin_dashboard_stats():
    db = SessionLocal()
    try:
        users = db.query(User).all()
        total_accounts = len(users)
        active_count = 0
        expiring_soon_count = 0
        expired_count = 0
        lifetime_count = 0

        for u in users:
            info = u.get_license_info()
            if not u.is_active:
                continue
            if info['status'] == 'expired':
                expired_count += 1
            elif info['status'] == 'expiring_soon':
                expiring_soon_count += 1
                active_count += 1
            elif info['status'] == 'lifetime':
                lifetime_count += 1
                active_count += 1
            else:
                active_count += 1

        total_customers = db.query(Customer).count()

        return jsonify({
            'total_users': total_accounts,
            'active_licenses': active_count,
            'expiring_soon': expiring_soon_count,
            'expired_licenses': expired_count,
            'lifetime_licenses': lifetime_count,
            'total_customers_in_db': total_customers
        })
    finally:
        db.close()

@app.route('/api/users', methods=['GET'])
@admin_required
def list_users():
    db = SessionLocal()
    try:
        users = db.query(User).order_by(User.id.asc()).all()
        return jsonify([u.to_dict() for u in users])
    finally:
        db.close()

@app.route('/api/users', methods=['POST'])
@admin_required
def create_new_user():
    data = request.json or {}
    username = data.get('username', '').strip().lower()
    password = data.get('password', '').strip()
    agency_name = data.get('agency_name', '').strip()
    full_name = data.get('full_name', '').strip()
    role = data.get('role', 'staff').strip().lower()
    phone = data.get('phone', '').strip()
    valid_until = data.get('valid_until', '').strip() or None
    max_customers = int(data.get('max_customers', 25000))
    notes = data.get('notes', '').strip()

    if not username:
        return jsonify({'error': 'Username is required.'}), 400
    if len(username) < 3:
        return jsonify({'error': 'Username must be at least 3 characters.'}), 400
    if not password or len(password) < 4:
        return jsonify({'error': 'Password must be at least 4 characters.'}), 400
    if role not in ('admin', 'staff'):
        role = 'staff'

    db = SessionLocal()
    try:
        existing = db.query(User).filter_by(username=username).first()
        if existing:
            return jsonify({'error': f"Username '{username}' is already taken. Please choose another."}), 400

        user = User(
            username=username,
            full_name=full_name,
            agency_name=agency_name,
            role=role,
            phone=phone,
            valid_until=valid_until,
            max_customers=max_customers,
            notes=notes,
            is_active=True
        )
        user.set_password(password)
        db.add(user)
        db.commit()
        return jsonify({
            'status': 'success',
            'message': f"New client '{username}' created with license successfully!",
            'user': user.to_dict()
        })
    except Exception as e:
        db.rollback()
        return jsonify({'error': f'Failed to create user: {str(e)}'}), 500
    finally:
        db.close()

@app.route('/api/users/<int:uid>', methods=['PUT'])
@admin_required
def update_user_info(uid):
    data = request.json or {}
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(id=uid).first()
        if not user:
            return jsonify({'error': 'User not found.'}), 404

        if 'agency_name' in data:
            user.agency_name = data['agency_name'].strip()

        if 'full_name' in data:
            user.full_name = data['full_name'].strip()

        if 'phone' in data:
            user.phone = data['phone'].strip()

        if 'notes' in data:
            user.notes = data['notes'].strip()

        if 'valid_until' in data:
            val = data['valid_until']
            user.valid_until = val.strip() if val else None

        if 'max_customers' in data:
            try:
                user.max_customers = int(data['max_customers'])
            except Exception:
                pass

        role = data.get('role')
        if role and role in ('admin', 'staff'):
            if user.id == flask_session.get('user_id') and role != 'admin':
                return jsonify({'error': 'You cannot remove administrative privileges from your own active account.'}), 400
            user.role = role

        password = data.get('password', '').strip()
        if password:
            if len(password) < 4:
                return jsonify({'error': 'Password must be at least 4 characters.'}), 400
            user.set_password(password)

        if 'is_active' in data:
            if user.id == flask_session.get('user_id') and not data['is_active']:
                return jsonify({'error': 'You cannot deactivate your own active session.'}), 400
            user.is_active = bool(data['is_active'])

        db.commit()
        return jsonify({
            'status': 'success',
            'message': f"Client '{user.username}' updated successfully!",
            'user': user.to_dict()
        })
    except Exception as e:
        db.rollback()
        return jsonify({'error': f'Failed to update user: {str(e)}'}), 500
    finally:
        db.close()

@app.route('/api/users/<int:uid>/renew', methods=['POST'])
@admin_required
def renew_user_license(uid):
    data = request.json or {}
    days = data.get('days') # e.g. 30, 90, 180, 365
    new_date = data.get('valid_until') # e.g. '2027-10-15'

    db = SessionLocal()
    try:
        user = db.query(User).filter_by(id=uid).first()
        if not user:
            return jsonify({'error': 'User not found.'}), 404

        today = datetime.date.today()
        if new_date:
            user.valid_until = new_date.strip()
        elif days is not None:
            try:
                days_int = int(days)
                # If current validity is active in the future, extend from that date; else extend from today
                base_date = today
                if user.valid_until:
                    try:
                        parsed = datetime.datetime.strptime(user.valid_until.strip(), '%Y-%m-%d').date()
                        if parsed >= today:
                            base_date = parsed
                    except Exception:
                        pass

                new_expiry = base_date + datetime.timedelta(days=days_int)
                user.valid_until = new_expiry.strftime('%Y-%m-%d')
            except Exception as e:
                return jsonify({'error': f'Invalid days count: {e}'}), 400
        else:
            return jsonify({'error': 'Either days or valid_until is required.'}), 400

        user.is_active = True # Re-activate if was suspended
        db.commit()
        return jsonify({
            'status': 'success',
            'message': f"लायसन्स यशस्वीरित्या वाढवले! नवीन अंतिम तारीख: {user.valid_until}",
            'user': user.to_dict()
        })
    except Exception as e:
        db.rollback()
        return jsonify({'error': f'Failed to renew license: {str(e)}'}), 500
    finally:
        db.close()

@app.route('/api/users/<int:uid>', methods=['DELETE'])
@admin_required
def delete_user_account(uid):
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(id=uid).first()
        if not user:
            return jsonify({'error': 'User not found.'}), 404

        if user.id == flask_session.get('user_id'):
            return jsonify({'error': 'You cannot delete your own active account.'}), 400
        if user.username == 'admin':
            return jsonify({'error': 'The master admin account cannot be deleted.'}), 400

        db.delete(user)
        db.commit()
        return jsonify({'status': 'success', 'message': f"User account '{user.username}' deleted."})
    except Exception as e:
        db.rollback()
        return jsonify({'error': f'Failed to delete user: {str(e)}'}), 500
    finally:
        db.close()

@app.route('/api/users/<int:uid>/toggle-status', methods=['POST'])
@admin_required
def toggle_user_active_status(uid):
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(id=uid).first()
        if not user:
            return jsonify({'error': 'User not found.'}), 404

        if user.id == flask_session.get('user_id'):
            return jsonify({'error': 'You cannot deactivate your own active account.'}), 400

        user.is_active = not user.is_active
        db.commit()
        status_word = 'Active' if user.is_active else 'Deactivated'
        return jsonify({
            'status': 'success',
            'is_active': user.is_active,
            'message': f"User '{user.username}' is now {status_word}."
        })
    except Exception as e:
        db.rollback()
        return jsonify({'error': f'Failed to toggle status: {str(e)}'}), 500
    finally:
        db.close()

# --- License Activation Routes ---
@app.route('/activate', methods=['GET'])
def activate_page():
    activated, _, _ = is_software_activated()
    if activated:
        return redirect(url_for('index'))
    return render_template('activate.html', machine_id=get_machine_id())

@app.route('/api/license/machine-id', methods=['GET'])
def api_get_machine_id():
    return jsonify({'machine_id': get_machine_id()})

@app.route('/api/license/activate', methods=['POST'])
def api_activate_license():
    data = request.json or {}
    key = data.get('license_key', '').strip()
    if not key:
        return jsonify({'status': 'error', 'error': 'License key is required.'}), 400

    valid, msg, info = verify_license_token(key)
    if not valid:
        return jsonify({'status': 'error', 'error': msg}), 400

    save_license(key)
    return jsonify({
        'status': 'success',
        'message': 'Software activated successfully!',
        'client_name': info.get('client_name', 'Client'),
        'expiry': info.get('expiry', 'LIFETIME'),
        'machine_id': info.get('machine_id', '')
    })

@app.route('/api/license/status', methods=['GET'])
def api_license_status():
    details = get_license_details()
    return jsonify({
        'machine_id': get_machine_id(),
        'details': details or {'is_valid': False, 'status_text': 'Unactivated'}
    })

# --- Health Check Route for Cloud Load Balancers ---
@app.route('/health')
@app.route('/healthz')
def health_check():
    return jsonify({'status': 'healthy', 'app': 'PassWish', 'timestamp': datetime.datetime.utcnow().isoformat()}), 200

# --- Web UI Routes ---
@app.route('/')
@login_required
def index():
    user_info = {
        'id': flask_session.get('user_id'),
        'username': flask_session.get('username'),
        'role': flask_session.get('role', 'staff'),
        'full_name': flask_session.get('full_name', '')
    }
    return render_template('index.html', current_user=user_info)

# --- Profile & Password API ---
@app.route('/api/profile', methods=['GET', 'POST'])
@login_required
def user_profile():
    db = SessionLocal()
    try:
        user_id = flask_session.get('user_id')
        user = db.query(User).filter_by(id=user_id).first()
        if not user:
            return jsonify({'error': 'User not found'}), 404

        if request.method == 'GET':
            def get_u_val(k, fallback=''):
                s = db.query(Setting).filter_by(key=f"user_{user_id}_{k}").first()
                if s and s.value:
                    return s.value
                return fallback

            return jsonify({
                'username': user.username,
                'shop_name': get_u_val('shop_name', user.agency_name or ''),
                'owner_name': get_u_val('owner_name', user.full_name or ''),
                'phone': get_u_val('phone', user.phone or ''),
                'address': get_u_val('address', '')
            })

        data = request.json or {}
        shop_name = data.get('shop_name', '').strip()
        owner_name = data.get('owner_name', '').strip()
        phone = data.get('phone', '').strip()
        address = data.get('address', '').strip()

        # Update User model
        user.agency_name = shop_name
        user.full_name = owner_name
        user.phone = phone

        # Save user-scoped settings
        def set_user_val(k, v):
            scoped_key = f"user_{user_id}_{k}"
            stg = db.query(Setting).filter_by(key=scoped_key).first()
            if stg:
                stg.value = v
            else:
                db.add(Setting(key=scoped_key, value=v))

        set_user_val('shop_name', shop_name)
        set_user_val('owner_name', owner_name)
        set_user_val('phone', phone)
        set_user_val('address', address)

        db.commit()
        return jsonify({'status': 'success', 'message': 'Profile and shop details saved successfully.'})
    finally:
        db.close()

@app.route('/api/change-password', methods=['POST'])
@login_required
def change_password():
    db = SessionLocal()
    try:
        user_id = flask_session.get('user_id')
        user = db.query(User).filter_by(id=user_id).first()
        if not user:
            return jsonify({'error': 'User not found'}), 404

        data = request.json or {}
        current_password = data.get('current_password', '').strip()
        new_password = data.get('new_password', '').strip()

        if not current_password or not new_password:
            return jsonify({'error': 'Current and new password are required.'}), 400

        if not user.check_password(current_password):
            return jsonify({'error': 'Current password is incorrect.'}), 400

        if len(new_password) < 4:
            return jsonify({'error': 'New password must be at least 4 characters long.'}), 400

        user.set_password(new_password)
        db.commit()
        return jsonify({'status': 'success', 'message': 'Password updated successfully.'})
    finally:
        db.close()

# --- API Routes ---

@app.route('/api/stats', methods=['GET'])
@login_required
def get_stats():
    session = SessionLocal()
    try:
        base_query = get_tenant_query(session, Customer)
        total_customers = base_query.count()

        today = datetime.date.today()
        today_md = today.strftime('%m-%d')
        tomorrow = today + datetime.timedelta(days=1)
        tomorrow_md = tomorrow.strftime('%m-%d')

        # Birthdays today
        birthdays_today = base_query.filter(
            Customer.dob.like(f"%-{today_md}")
        ).count()

        # Birthdays this week (next 7 days)
        week_mds = [(today + datetime.timedelta(days=i)).strftime('%m-%d') for i in range(7)]
        bday_conditions = [Customer.dob.like(f"%-{md}") for md in week_mds]
        birthdays_week = base_query.filter(or_(*bday_conditions)).count() if bday_conditions else 0

        # Passport expiries in next 6 months
        six_months_ahead = (today + datetime.timedelta(days=180)).strftime('%Y-%m-%d')
        today_str = today.strftime('%Y-%m-%d')
        expiries_6m = base_query.filter(
            and_(Customer.expiry_date >= today_str, Customer.expiry_date <= six_months_ahead)
        ).count()

        templates_count = session.query(MessageTemplate).count()
        recent_backups = len(list_backups())

        db_type = "Cloud PostgreSQL (Neon)"

        return jsonify({
            'total_customers': total_customers,
            'birthdays_today': birthdays_today,
            'birthdays_week': birthdays_week,
            'expiries_6m': expiries_6m,
            'templates_count': templates_count,
            'recent_backups': recent_backups,
            'db_type': db_type,
            'today_date': today.strftime('%d %B %Y'),
            'machine_id': get_machine_id(),
            'license': get_license_details()
        })
    finally:
        session.close()

@app.route('/api/customers', methods=['GET'])
@login_required
def get_customers():
    session = SessionLocal()
    try:
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 25))
        search = request.args.get('q', '').strip()
        month = request.args.get('month', '').strip()
        expiry_window = request.args.get('expiry_window', '').strip()

        query = get_tenant_query(session, Customer)

        if search:
            search_pattern = f"%{search}%"
            query = query.filter(
                or_(
                    Customer.name.ilike(search_pattern),
                    Customer.phone.ilike(search_pattern),
                    Customer.phone2.ilike(search_pattern),
                    Customer.passport_no.ilike(search_pattern),
                    Customer.city.ilike(search_pattern)
                )
            )

        if month:
            # Match MM in YYYY-MM-DD
            month_fmt = f"-{int(month):02d}-"
            query = query.filter(Customer.dob.like(f"%{month_fmt}%"))

        if expiry_window:
            today_str = datetime.date.today().strftime('%Y-%m-%d')
            days = 90 if expiry_window == '3m' else (180 if expiry_window == '6m' else 365)
            end_date_str = (datetime.date.today() + datetime.timedelta(days=days)).strftime('%Y-%m-%d')
            query = query.filter(and_(Customer.expiry_date >= today_str, Customer.expiry_date <= end_date_str))

        service_status = request.args.get('service_status', '').strip()
        if service_status and service_status != 'all':
            query = query.filter(Customer.service_status == service_status)

        wa_status = request.args.get('wa_status', '').strip()
        if wa_status and wa_status != 'all':
            query = query.filter(Customer.whatsapp_status == wa_status)

        total = query.count()
        customers = query.order_by(Customer.id.desc()).offset((page - 1) * limit).limit(limit).all()

        return jsonify({
            'customers': [c.to_dict() for c in customers],
            'total': total,
            'page': page,
            'limit': limit,
            'pages': (total + limit - 1) // limit if limit else 1
        })
    finally:
        session.close()

@app.route('/api/customers/export', methods=['GET'])
@login_required
def export_customers_excel():
    session = SessionLocal()
    try:
        search = request.args.get('q', '').strip()
        month = request.args.get('month', '').strip()
        expiry_window = request.args.get('expiry_window', '').strip()
        service_status = request.args.get('service_status', '').strip()
        wa_status = request.args.get('wa_status', '').strip()

        query = get_tenant_query(session, Customer)

        if search:
            search_pattern = f"%{search}%"
            query = query.filter(
                or_(
                    Customer.name.ilike(search_pattern),
                    Customer.phone.ilike(search_pattern),
                    Customer.phone2.ilike(search_pattern),
                    Customer.passport_no.ilike(search_pattern),
                    Customer.city.ilike(search_pattern)
                )
            )

        if month:
            month_fmt = f"-{int(month):02d}-"
            query = query.filter(Customer.dob.like(f"%{month_fmt}%"))

        if expiry_window:
            today_str = datetime.date.today().strftime('%Y-%m-%d')
            days = 90 if expiry_window == '3m' else (180 if expiry_window == '6m' else 365)
            end_date_str = (datetime.date.today() + datetime.timedelta(days=days)).strftime('%Y-%m-%d')
            query = query.filter(and_(Customer.expiry_date >= today_str, Customer.expiry_date <= end_date_str))

        if service_status and service_status != 'all':
            query = query.filter(Customer.service_status == service_status)

        if wa_status and wa_status != 'all':
            query = query.filter(Customer.whatsapp_status == wa_status)

        customers = query.order_by(Customer.id.desc()).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customers"

        headers = ["Sr No", "Customer Name", "Primary Mobile", "Secondary Mobile", "Date of Birth", "Passport No", "Issue Date", "Expiry Date", "City", "Pipeline Stage", "WhatsApp Status", "Notes"]
        ws.append(headers)

        header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="E4E4E7"),
            right=Side(style="thin", color="E4E4E7"),
            top=Side(style="thin", color="E4E4E7"),
            bottom=Side(style="thin", color="E4E4E7")
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        for idx, c in enumerate(customers, 1):
            row_vals = [
                idx,
                c.name or '',
                c.phone or '',
                c.phone2 or '',
                c.dob or '',
                c.passport_no or '',
                c.issue_date or '',
                c.expiry_date or '',
                c.city or '',
                c.service_status or 'New',
                c.whatsapp_status or 'active',
                c.notes or ''
            ]
            ws.append(row_vals)
            row_idx = ws.max_row
            ws.row_dimensions[row_idx].height = 20
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if col_idx in (1, 3, 4, 5, 7, 8, 10, 11):
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        col_widths = [8, 24, 16, 16, 14, 16, 14, 14, 16, 18, 16, 30]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        today_str = datetime.date.today().strftime('%Y%m%d')
        out_path = os.path.join(app.config['UPLOAD_FOLDER'], f"PassWish_Customers_{today_str}.xlsx")
        wb.save(out_path)
        return send_file(out_path, as_attachment=True, download_name=os.path.basename(out_path), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        session.close()

@app.route('/api/customers', methods=['POST'])
@login_required
def add_customer():
    session = SessionLocal()
    try:
        data = request.json or {}
        name = data.get('name', '').strip()
        raw_phone = data.get('phone', '').strip()
        raw_phone2 = data.get('phone2', '').strip()

        if not name or not raw_phone:
            return jsonify({'error': 'Customer Name and Primary Mobile Number are required.'}), 400

        phone = clean_phone(raw_phone)
        if not phone or len(phone) < 10:
            return jsonify({'error': 'Please enter a valid 10-digit primary mobile number.'}), 400

        phone2 = clean_phone(raw_phone2) if raw_phone2 else ''

        # Check existing phone in THIS shop/tenant
        uid = flask_session.get('user_id')
        role = flask_session.get('role', 'staff')
        exist_q = session.query(Customer).filter_by(phone=phone)
        if role != 'admin':
            exist_q = exist_q.filter_by(created_by_user_id=uid)
        existing = exist_q.first()
        if existing:
            return jsonify({'error': f'Customer with primary phone number ({phone}) already exists in your list: {existing.name}.'}), 400

        c = Customer(
            name=name,
            phone=phone,
            phone2=phone2,
            dob=data.get('dob', '').strip() if data.get('dob') else '',
            passport_no=data.get('passport_no', '').strip() if data.get('passport_no') else '',
            expiry_date=data.get('expiry_date', '').strip() if data.get('expiry_date') else '',
            issue_date=data.get('issue_date', '').strip() if data.get('issue_date') else '',
            city=data.get('city', '').strip() if data.get('city') else '',
            notes=data.get('notes', '').strip() if data.get('notes') else '',
            created_by_user_id=uid
        )
        session.add(c)
        session.commit()
        return jsonify({'status': 'success', 'customer': c.to_dict()})
    except Exception as e:
        session.rollback()
        return jsonify({'error': f'Failed to add customer: {str(e)}'}), 500
    finally:
        session.close()

@app.route('/api/customers/<int:cid>', methods=['PUT', 'DELETE'])
@login_required
def manage_customer(cid):
    session = SessionLocal()
    try:
        c = session.query(Customer).filter_by(id=cid).first()
        if not c:
            return jsonify({'error': 'Customer not found'}), 404

        uid = flask_session.get('user_id')
        role = flask_session.get('role', 'staff')
        if role != 'admin' and c.created_by_user_id != uid:
            return jsonify({'error': 'Access denied: You do not have permission to modify this customer.'}), 403

        if request.method == 'DELETE':
            session.delete(c)
            session.commit()
            return jsonify({'status': 'success', 'message': 'Customer deleted successfully.'})

        data = request.json or {}
        c.name = data.get('name', c.name).strip()
        raw_phone = data.get('phone', '').strip()
        if raw_phone:
            cleaned = clean_phone(raw_phone)
            if cleaned:
                c.phone = cleaned
        raw_phone2 = data.get('phone2')
        if raw_phone2 is not None:
            c.phone2 = clean_phone(raw_phone2) if raw_phone2.strip() else ''
        if data.get('dob') is not None:
            c.dob = data.get('dob', '').strip()
        if data.get('passport_no') is not None:
            c.passport_no = data.get('passport_no', '').strip()
        if data.get('expiry_date') is not None:
            c.expiry_date = data.get('expiry_date', '').strip()
        if data.get('issue_date') is not None:
            c.issue_date = data.get('issue_date', '').strip()
        if data.get('city') is not None:
            c.city = data.get('city', '').strip()
        if data.get('notes') is not None:
            c.notes = data.get('notes', '').strip()
        if data.get('whatsapp_status') is not None:
            c.whatsapp_status = data.get('whatsapp_status', c.whatsapp_status).strip()
        session.commit()
        return jsonify({'status': 'success', 'customer': c.to_dict()})
    except Exception as e:
        session.rollback()
        return jsonify({'error': f'Failed to update customer: {str(e)}'}), 500
    finally:
        session.close()

@app.route('/api/customers/<int:cid>/toggle-whatsapp', methods=['POST'])
@login_required
def toggle_customer_whatsapp(cid):
    session = SessionLocal()
    try:
        c = session.query(Customer).filter_by(id=cid).first()
        if not c:
            return jsonify({'error': 'Customer not found'}), 404

        uid = flask_session.get('user_id')
        role = flask_session.get('role', 'staff')
        if role != 'admin' and c.created_by_user_id != uid:
            return jsonify({'error': 'Access denied: Unauthorized'}), 403

        data = request.json or {}
        new_status = data.get('status')
        if not new_status:
            new_status = 'no_whatsapp' if c.whatsapp_status == 'active' else 'active'
        c.whatsapp_status = new_status
        session.commit()
        return jsonify({'status': 'success', 'whatsapp_status': c.whatsapp_status, 'customer': c.to_dict()})
    finally:
        session.close()

@app.route('/api/customers/<int:cid>/service-status', methods=['POST'])
@login_required
def update_customer_service_status(cid):
    session = SessionLocal()
    try:
        c = session.query(Customer).filter_by(id=cid).first()
        if not c:
            return jsonify({'error': 'Customer not found'}), 404

        uid = flask_session.get('user_id')
        role = flask_session.get('role', 'staff')
        if role != 'admin' and c.created_by_user_id != uid:
            return jsonify({'error': 'Access denied: Unauthorized'}), 403

        data = request.json or {}
        new_status = data.get('status', 'New').strip()
        c.service_status = new_status
        session.commit()
        return jsonify({'status': 'success', 'service_status': c.service_status, 'customer': c.to_dict()})
    finally:
        session.close()

@app.route('/uploads/<path:filename>')
def serve_uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/api/upload/poster', methods=['POST'])
@login_required
def upload_poster():
    try:
        if 'image' not in request.files:
            return jsonify({'error': 'No image file provided'}), 400
        f = request.files['image']
        if not f.filename:
            return jsonify({'error': 'Empty filename'}), 400
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ('.jpg', '.jpeg', '.png', '.webp'):
            return jsonify({'error': 'Only JPG, PNG and WebP images are supported'}), 400
        
        safe_name = f"poster_{int(time.time()*1000)}{ext}"
        save_path = os.path.join(app.config['UPLOAD_FOLDER'], safe_name)
        f.save(save_path)
        
        return jsonify({
            'status': 'success',
            'filename': safe_name,
            'image_url': f"/uploads/{safe_name}",
            'url': f"/uploads/{safe_name}",
            'file_path': os.path.abspath(save_path),
            'server_path': os.path.abspath(save_path)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/reminder/test', methods=['POST'])
@login_required
def test_reminder_notification():
    try:
        trigger_desktop_notification("PassWish — Test Notification", "Morning reminder notifications are working perfectly on your Windows desktop!")
        return jsonify({'status': 'success', 'message': 'Notification sent to desktop'})
    except Exception as e:
        return jsonify({'status': 'error', 'error': str(e)}), 500

@app.route('/api/template/sample-excel', methods=['GET'])
def download_sample_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PassWish_Sample"

    headers = [
        "Customer Name", "Primary Mobile", "Secondary Mobile",
        "Date of Birth (YYYY-MM-DD)", "Passport Number",
        "Issue Date (YYYY-MM-DD)", "Expiry Date (YYYY-MM-DD)", "City", "Notes"
    ]
    ws.append(headers)

    sample_rows = [
        ["Rahul Sharma", "9822012345", "9822012346", "1992-05-14", "V8921453", "2018-02-10", "2028-02-09", "Pune", "Urgent Renewal"],
        ["Pooja Patil", "9890123456", "", "1995-10-22", "M4512987", "2016-08-15", "2026-08-14", "Kolhapur", "Tatkal Client"]
    ]
    for r in sample_rows:
        ws.append(r)

    header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    col_widths = [22, 18, 18, 25, 18, 24, 24, 16, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    out_path = os.path.join(app.config['UPLOAD_FOLDER'], "PassWish_Sample_Customer_Template.xlsx")
    wb.save(out_path)
    return send_file(out_path, as_attachment=True, download_name="PassWish_Sample_Customer_Template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# --- Excel Upload & Smart Weekly Append ---
@app.route('/api/upload/preview', methods=['POST'])
@login_required
def upload_preview():
    try:
        if 'file' not in request.files:
            return jsonify({'status': 'error', 'message': 'No file uploaded'}), 400
        file = request.files['file']
        if not file.filename:
            return jsonify({'status': 'error', 'message': 'No file selected'}), 400

        filename = secure_filename(file.filename) or 'upload_file.xlsx'
        save_path = os.path.join(app.config['UPLOAD_FOLDER'], f"preview_{filename}")
        file.save(save_path)

        result = get_preview_data(save_path)
        return jsonify(result)
    except Exception as e:
        return jsonify({'status': 'error', 'message': f'File read error: {str(e)}'}), 500

@app.route('/api/upload/process', methods=['POST'])
@login_required
def upload_process():
    try:
        if 'file' not in request.files:
            return jsonify({'status': 'error', 'message': 'No file uploaded'}), 400
        file = request.files['file']
        mode = request.form.get('mode', 'append_new') # append_new, update_existing, replace
        custom_mapping_raw = request.form.get('mapping')
        custom_mapping = json.loads(custom_mapping_raw) if custom_mapping_raw else None

        filename = secure_filename(file.filename) or 'upload_file.xlsx'
        save_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(save_path)

        result = process_excel_file(save_path, mode=mode, custom_mapping=custom_mapping, user_id=flask_session.get('user_id'))
        return jsonify(result)
    except Exception as e:
        return jsonify({'status': 'error', 'message': f'Import error: {str(e)}'}), 500

# --- Birthday Tracker ---
@app.route('/api/birthdays', methods=['GET'])
@login_required
def get_birthdays():
    session = SessionLocal()
    try:
        timeframe = request.args.get('timeframe', 'today') # today, tomorrow, week
        template_id = request.args.get('template_id')

        # Get default birthday template
        tpl_query = session.query(MessageTemplate).filter_by(category='birthday')
        if template_id:
            tpl = session.query(MessageTemplate).filter_by(id=template_id).first()
        else:
            tpl = tpl_query.filter_by(is_default=True).first() or tpl_query.first()

        body_tpl = tpl.body if tpl else 'Dear {name}, wishing you a very Happy Birthday! 🎂'

        today = datetime.date.today()
        if timeframe == 'today':
            target_dates = [today.strftime('%m-%d')]
        elif timeframe == 'tomorrow':
            target_dates = [(today + datetime.timedelta(days=1)).strftime('%m-%d')]
        else: # week
            target_dates = [(today + datetime.timedelta(days=i)).strftime('%m-%d') for i in range(7)]

        base_query = get_tenant_query(session, Customer)
        conditions = [Customer.dob.like(f"%-{md}") for md in target_dates]
        customers = base_query.filter(or_(*conditions)).all() if conditions else []

        results = []
        for c in customers:
            msg = render_template_text(body_tpl, c)
            results.append({
                'customer': c.to_dict(),
                'rendered_message': msg,
                'whatsapp_url': build_whatsapp_url(c.phone, msg),
                'whatsapp_url_2': build_whatsapp_url(c.phone2, msg) if c.phone2 else ''
            })

        return jsonify({
            'timeframe': timeframe,
            'total': len(results),
            'birthdays': results,
            'template_used': tpl.to_dict() if tpl else None
        })
    finally:
        session.close()

# --- Festival / Event Broadcast ---
@app.route('/api/broadcast/preview', methods=['POST'])
@login_required
def preview_broadcast():
    session = SessionLocal()
    try:
        data = request.json or {}
        target = data.get('target', 'all') # all, birthday_today, expiring_6m, custom_city
        city_filter = data.get('city', '').strip()
        festival_name = data.get('festival_name', 'Festival')
        custom_body = data.get('body', '')
        template_id = data.get('template_id')

        if not custom_body and template_id:
            tpl = session.query(MessageTemplate).filter_by(id=template_id).first()
            if tpl:
                custom_body = tpl.body
        
        if not custom_body:
            # Fallback
            tpl = session.query(MessageTemplate).filter_by(category='diwali', is_default=True).first()
            custom_body = tpl.body if tpl else 'Happy Festival {name}! - {shop_name}'

        query = get_tenant_query(session, Customer)
        today = datetime.date.today()

        if target == 'birthday_today':
            today_md = today.strftime('%m-%d')
            query = query.filter(Customer.dob.like(f"%-{today_md}"))
        elif target == 'expiring_6m':
            today_str = today.strftime('%Y-%m-%d')
            end_date = (today + datetime.timedelta(days=180)).strftime('%Y-%m-%d')
            query = query.filter(and_(Customer.expiry_date >= today_str, Customer.expiry_date <= end_date))
        elif target == 'city' and city_filter:
            query = query.filter(Customer.city.ilike(f"%{city_filter}%"))

        customers = query.all()

        messages = []
        for c in customers:
            msg = render_template_text(custom_body, c, festival_name=festival_name)
            messages.append({
                'id': c.id,
                'name': c.name,
                'phone': c.phone,
                'phone2': c.phone2 or '',
                'city': c.city,
                'whatsapp_status': c.whatsapp_status or 'active',
                'message': msg,
                'whatsapp_url': build_whatsapp_url(c.phone, msg),
                'whatsapp_url_2': build_whatsapp_url(c.phone2, msg) if c.phone2 else ''
            })

        return jsonify({
            'total_recipients': len(messages),
            'festival_name': festival_name,
            'recipients': messages[:500], # return first 500 in UI preview
            'has_more': len(messages) > 500
        })
    finally:
        session.close()

# --- Passport Expiry Tracker ---
@app.route('/api/expiries', methods=['GET'])
@login_required
def get_expiries():
    session = SessionLocal()
    try:
        months = int(request.args.get('months', 6))
        template_id = request.args.get('template_id')

        tpl_query = session.query(MessageTemplate).filter_by(category='expiry')
        if template_id:
            tpl = session.query(MessageTemplate).filter_by(id=template_id).first()
        else:
            tpl = tpl_query.filter_by(is_default=True).first() or tpl_query.first()

        body_tpl = tpl.body if tpl else 'Passport Renewal Reminder for {name}. - {shop_name}'

        today = datetime.date.today()
        today_str = today.strftime('%Y-%m-%d')
        end_date = (today + datetime.timedelta(days=months * 30)).strftime('%Y-%m-%d')

        customers = get_tenant_query(session, Customer).filter(
            and_(Customer.expiry_date >= today_str, Customer.expiry_date <= end_date)
        ).order_by(Customer.expiry_date.asc()).all()

        results = []
        for c in customers:
            msg = render_template_text(body_tpl, c)
            results.append({
                'customer': c.to_dict(),
                'rendered_message': msg,
                'whatsapp_url': build_whatsapp_url(c.phone, msg),
                'whatsapp_url_2': build_whatsapp_url(c.phone2, msg) if c.phone2 else ''
            })

        return jsonify({
            'months': months,
            'total': len(results),
            'customers': results
        })
    finally:
        session.close()

# --- Templates CRUD ---
@app.route('/api/templates', methods=['GET', 'POST'])
@login_required
def manage_templates():
    session = SessionLocal()
    try:
        if request.method == 'GET':
            category = request.args.get('category')
            query = session.query(MessageTemplate)
            if category:
                query = query.filter_by(category=category)
            templates = query.order_by(MessageTemplate.id.asc()).all()
            return jsonify({'templates': [t.to_dict() for t in templates]})

        data = request.json or {}
        name = data.get('name', '').strip()
        body = data.get('body', '').strip()
        category = data.get('category', 'general').strip()
        if not name or not body:
            return jsonify({'error': 'Name and message body are required'}), 400

        tpl = MessageTemplate(
            name=name,
            body=body,
            category=category,
            is_default=bool(data.get('is_default', False))
        )
        session.add(tpl)
        session.commit()
        return jsonify({'status': 'success', 'template': tpl.to_dict()})
    finally:
        session.close()

@app.route('/api/templates/<int:tid>', methods=['PUT', 'DELETE'])
@login_required
def modify_template(tid):
    session = SessionLocal()
    try:
        tpl = session.query(MessageTemplate).filter_by(id=tid).first()
        if not tpl:
            return jsonify({'error': 'Template not found'}), 404

        if request.method == 'DELETE':
            session.delete(tpl)
            session.commit()
            return jsonify({'status': 'success', 'message': 'Template deleted successfully'})

        data = request.json or {}
        tpl.name = data.get('name', tpl.name)
        tpl.body = data.get('body', tpl.body)
        tpl.category = data.get('category', tpl.category)
        if 'is_default' in data:
            if data['is_default']:
                # Unset other defaults in same category
                session.query(MessageTemplate).filter_by(category=tpl.category).update({'is_default': False})
            tpl.is_default = bool(data['is_default'])

        session.commit()
        return jsonify({'status': 'success', 'template': tpl.to_dict()})
    finally:
        session.close()

# --- Log Sent Message ---
@app.route('/api/broadcast/log-send', methods=['POST'])
@login_required
def log_send():
    session = SessionLocal()
    try:
        data = request.json or {}
        log = SentLog(
            customer_id=data.get('customer_id'),
            customer_name=data.get('customer_name'),
            phone=data.get('phone', ''),
            message_type=data.get('message_type', 'whatsapp'),
            message_content=data.get('message_content', ''),
            status=data.get('status', 'sent'),
            sent_by_user_id=flask_session.get('user_id')
        )
        session.add(log)
        session.commit()
        return jsonify({'status': 'success', 'log_id': log.id})
    finally:
        session.close()

# --- Web-Native WhatsApp Dispatch Queue API (Zero Client Dependencies) ---
@app.route('/api/queue/prepare', methods=['POST'])
@login_required
def prepare_web_queue():
    """
    Prepares a list of WhatsApp messages for Web-Native browser dispatch.
    Renders variables ({name}, {passport_no}, {expiry_date}, {shop_name}, {shop_phone}, {shop_address})
    and returns direct web.whatsapp.com and api.whatsapp.com links.
    Zero client software/dependencies needed!
    """
    session = SessionLocal()
    try:
        data = request.json or {}
        qtype = data.get('type', 'custom')
        items = data.get('items', [])
        template_id = data.get('template_id')
        custom_text = data.get('custom_text', '')
        festival_name = data.get('festival_name', '')

        # Resolve template if not provided
        body_template = custom_text
        if not body_template and template_id:
            tpl = session.query(MessageTemplate).filter_by(id=template_id).first()
            if tpl:
                body_template = tpl.body

        if not body_template:
            cat = 'birthday' if qtype == 'birthday' else ('expiry' if qtype == 'expiry' else 'festival')
            tpl = session.query(MessageTemplate).filter_by(category=cat, is_default=True).first() or \
                  session.query(MessageTemplate).filter_by(category=cat).first()
            if tpl:
                body_template = tpl.body
            else:
                body_template = "Dear {name}, greetings from {shop_name}!"

        shop_name = get_setting('shop_name', '').strip()
        shop_phone = get_setting('phone', '').strip()
        shop_address = get_setting('address', '').strip()
        send_target = data.get('send_target', 'primary') # 'primary', 'secondary', 'both'

        prepared = []
        for it in items:
            cid = it.get('id') or it.get('customer_id')
            cname = it.get('name') or it.get('customer_name') or ''
            cphone = it.get('phone') or ''
            cphone2 = it.get('phone2') or ''
            cpassport = it.get('passport_no') or ''
            cexpiry = it.get('expiry_date') or ''
            ccity = it.get('city') or ''

            if cid and (not cname or not cphone2):
                cust = get_tenant_query(session, Customer).filter_by(id=cid).first()
                if cust:
                    cname = cname or cust.name
                    cphone = cphone or cust.phone
                    cphone2 = cphone2 or (cust.phone2 or '')
                    cpassport = cpassport or (cust.passport_no or '')
                    cexpiry = cexpiry or (cust.expiry_date or '')
                    ccity = ccity or (cust.city or '')

            phones_to_queue = []
            if send_target in ('primary', 'both') or not cphone2:
                phones_to_queue.append((cphone, f"{cname} (Primary)" if (send_target == 'both' and cphone2) else cname))
            if send_target in ('secondary', 'both') and cphone2:
                phones_to_queue.append((cphone2, f"{cname} (Alt Phone)"))

            for p_num, label_name in phones_to_queue:
                is_valid, clean_p, reason = validate_mobile_for_queue(p_num)

                # Render message text
                txt = body_template
                txt = txt.replace('{name}', cname)
                txt = txt.replace('{phone}', p_num)
                txt = txt.replace('{passport_no}', cpassport)
                txt = txt.replace('{expiry_date}', cexpiry)
                txt = txt.replace('{city}', ccity)
                txt = txt.replace('{shop_name}', shop_name)
                txt = txt.replace('{shop_phone}', shop_phone)
                txt = txt.replace('{shop_address}', shop_address)
                txt = txt.replace('{festival_name}', festival_name)

                encoded = urllib.parse.quote(txt)
                target_phone = clean_p if is_valid else re.sub(r'\D', '', str(p_num))
                if not target_phone.startswith('91') and len(target_phone) == 10:
                    target_phone = f"91{target_phone}"

                prepared.append({
                    'customer_id': cid,
                    'name': label_name,
                    'phone': p_num,
                    'clean_phone': target_phone,
                    'is_valid_mobile': is_valid,
                    'validation_note': reason,
                    'passport_no': cpassport,
                    'expiry_date': cexpiry,
                    'message_text': txt,
                    'web_url': f"https://web.whatsapp.com/send?phone={target_phone}&text={encoded}",
                    'mobile_url': f"https://api.whatsapp.com/send?phone={target_phone}&text={encoded}"
                })

        return jsonify({
            'status': 'success',
            'total': len(prepared),
            'items': prepared
        })
    finally:
        session.close()

# --- Reports & Logs CRUD & Export (Tenant Isolated) ---
@app.route('/api/reports/logs', methods=['GET'])
@login_required
def get_report_logs():
    session = SessionLocal()
    try:
        page = int(request.args.get('page', 1))
        limit = int(request.args.get('limit', 25))
        category = request.args.get('category', 'all').strip()
        status = request.args.get('status', 'all').strip()
        timeframe = request.args.get('timeframe', 'all').strip()
        search = request.args.get('q', '').strip()

        query = get_tenant_query(session, SentLog)

        if category and category != 'all':
            query = query.filter(SentLog.message_type == category)

        if status and status != 'all':
            query = query.filter(SentLog.status == status)

        today = datetime.date.today()
        if timeframe == 'today':
            start_today = datetime.datetime.combine(today, datetime.time.min)
            query = query.filter(SentLog.sent_at >= start_today)
        elif timeframe == 'week':
            start_week = datetime.datetime.combine(today - datetime.timedelta(days=7), datetime.time.min)
            query = query.filter(SentLog.sent_at >= start_week)
        elif timeframe == 'month':
            start_month = datetime.datetime.combine(today - datetime.timedelta(days=30), datetime.time.min)
            query = query.filter(SentLog.sent_at >= start_month)

        if search:
            sp = f"%{search}%"
            query = query.filter(or_(
                SentLog.customer_name.ilike(sp),
                SentLog.phone.ilike(sp),
                SentLog.message_content.ilike(sp)
            ))

        total = query.count()
        logs = query.order_by(SentLog.id.desc()).offset((page - 1) * limit).limit(limit).all()

        # Summary Statistics (Tenant Scoped)
        all_logs = get_tenant_query(session, SentLog)
        total_all = all_logs.count()
        total_sent = all_logs.filter_by(status='sent').count()
        total_skipped = all_logs.filter(SentLog.status.in_(['failed_no_whatsapp', 'skipped_invalid', 'skipped_no_whatsapp'])).count()
        total_bday = all_logs.filter_by(message_type='birthday').count()
        total_expiry = all_logs.filter_by(message_type='expiry').count()
        total_broadcast = all_logs.filter_by(message_type='broadcast').count()

        summary_dict = {
            'total': total_all,
            'sent': total_sent,
            'skipped': total_skipped,
            'birthday': total_bday,
            'expiry': total_expiry,
            'broadcast': total_broadcast
        }
        return jsonify({
            'logs': [l.to_dict() for l in logs],
            'total': total,
            'page': page,
            'limit': limit,
            'pages': (total + limit - 1) // limit if limit else 1,
            'stats': summary_dict,
            'summary': summary_dict
        })
    finally:
        session.close()

@app.route('/api/reports/export', methods=['GET'])
@login_required
def export_reports_excel():
    session = SessionLocal()
    try:
        category = request.args.get('category', 'all').strip()
        status = request.args.get('status', 'all').strip()
        timeframe = request.args.get('timeframe', 'all').strip()

        query = get_tenant_query(session, SentLog)
        if category and category != 'all':
            query = query.filter(SentLog.message_type == category)
        if status and status != 'all':
            query = query.filter(SentLog.status == status)

        today = datetime.date.today()
        if timeframe == 'today':
            query = query.filter(SentLog.sent_at >= datetime.datetime.combine(today, datetime.time.min))
        elif timeframe == 'week':
            query = query.filter(SentLog.sent_at >= datetime.datetime.combine(today - datetime.timedelta(days=7), datetime.time.min))
        elif timeframe == 'month':
            query = query.filter(SentLog.sent_at >= datetime.datetime.combine(today - datetime.timedelta(days=30), datetime.time.min))

        records = query.order_by(SentLog.id.desc()).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PassWish Sent History"

        headers = ["Sr No", "Date & Time", "Customer Name", "Phone", "Message Category", "Delivery Status", "Message Content"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        status_sent_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
        status_skip_fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
        status_font_sent = Font(name="Calibri", size=10, bold=True, color="2E7D32")
        status_font_skip = Font(name="Calibri", size=10, bold=True, color="E65100")

        for idx, r in enumerate(records, 1):
            dt_str = r.sent_at.strftime('%Y-%m-%d %H:%M:%S') if r.sent_at else ''
            st_text = 'Sent' if r.status == 'sent' else ('No WhatsApp / Skipped' if 'no_whatsapp' in (r.status or '') else r.status)
            row_vals = [idx, dt_str, r.customer_name or '', r.phone or '', (r.message_type or '').title(), st_text, (r.message_content or '')]
            ws.append(row_vals)
            row_idx = ws.max_row
            ws.row_dimensions[row_idx].height = 22

            for col_idx in range(1, len(row_vals) + 1):
                c_cell = ws.cell(row=row_idx, column=col_idx)
                c_cell.border = thin_border
                c_cell.font = Font(name="Calibri", size=10)
                if col_idx in (1, 2, 5, 6):
                    c_cell.alignment = center_align
                else:
                    c_cell.alignment = left_align

                if col_idx == 6:
                    if r.status == 'sent':
                        c_cell.fill = status_sent_fill
                        c_cell.font = status_font_sent
                    else:
                        c_cell.fill = status_skip_fill
                        c_cell.font = status_font_skip

        col_widths = [8, 20, 24, 18, 18, 22, 60]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        out_path = os.path.join(app.config['UPLOAD_FOLDER'], f"PassWish_Sent_Report_{today.strftime('%Y%m%d')}.xlsx")
        wb.save(out_path)
        return send_file(out_path, as_attachment=True, download_name=os.path.basename(out_path), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        session.close()

# --- Backup & Restore (Tenant Scoped) ---
@app.route('/api/backup/export', methods=['GET'])
@login_required
def export_backup():
    try:
        uid = flask_session.get('user_id')
        filepath, count, filename = create_backup(backup_type='manual_export', user_id=uid)
        return send_file(filepath, as_attachment=True, download_name=filename, mimetype='application/json')
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/backup/restore', methods=['POST'])
@login_required
def restore_db_backup():
    if 'file' not in request.files:
        return jsonify({'error': 'Backup file is required'}), 400
    file = request.files['file']
    mode = request.form.get('mode', 'merge') # merge or replace

    try:
        uid = flask_session.get('user_id')
        content = json.load(file)
        result = restore_backup(content, mode=mode, user_id=uid)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': f'Restore error: {str(e)}'}), 500

@app.route('/api/backup/snapshots', methods=['GET'])
@login_required
def get_snapshots():
    backups = list_backups()
    return jsonify({'snapshots': backups})

# --- Settings (Isolated per Shop/Tenant) ---
@app.route('/api/settings', methods=['GET', 'POST'])
@login_required
def manage_settings():
    session = SessionLocal()
    uid = flask_session.get('user_id')
    try:
        if request.method == 'GET':
            EXCLUDE_GLOBAL = {'shop_name', 'owner_name', 'phone', 'address'}
            # Global fallback settings (only operational configs like reminder/whatsapp)
            settings = {
                s.key: s.value for s in session.query(Setting).filter(~Setting.key.like('user_%')).all()
                if s.key not in EXCLUDE_GLOBAL
            }
            # Overlay user/shop specific settings
            if uid:
                user_prefix = f"user_{uid}_"
                user_settings = session.query(Setting).filter(Setting.key.like(f"{user_prefix}%")).all()
                for us in user_settings:
                    clean_k = us.key[len(user_prefix):]
                    settings[clean_k] = us.value

                # If shop_name is empty, default to user's agency_name
                if not settings.get('shop_name'):
                    u = session.query(User).filter_by(id=uid).first()
                    if u and u.agency_name:
                        settings['shop_name'] = u.agency_name
                    if u and u.phone and not settings.get('phone'):
                        settings['phone'] = u.phone
                    if u and u.full_name and not settings.get('owner_name'):
                        settings['owner_name'] = u.full_name
            return jsonify({'settings': settings})

        data = request.json or {}
        for k, v in data.items():
            full_key = f"user_{uid}_{k}" if uid else k
            stg = session.query(Setting).filter_by(key=full_key).first()
            if stg:
                stg.value = str(v)
            else:
                session.add(Setting(key=full_key, value=str(v)))
        session.commit()
        return jsonify({'status': 'success', 'message': 'Settings saved successfully!'})
    finally:
        session.close()

# --- Automated WhatsApp Bot Background Worker ---
bot_state = {
    'running': False,
    'stop_requested': False,
    'skip_requested': False,
    'type': '',
    'current': 0,
    'total': 0,
    'current_id': 0,
    'current_name': '',
    'current_phone': '',
    'current_message': '',
    'sent_count': 0,
    'skipped_count': 0,
    'skipped_items': [],
    'status': 'idle', # idle, starting, running, loading_whatsapp, sending, closing_tab, waiting_next, finished, stopped, error
    'status_text': 'Idle',
    'error': ''
}
bot_lock = threading.Lock()

def bot_worker_thread(queue_type, delay_sec=5, custom_items=None, global_image_path=None):
    global bot_state
    try:
        import pyautogui
        import pyperclip
        import webbrowser
        pyautogui.FAILSAFE = True
    except (ImportError, Exception) as ie:
        with bot_lock:
            bot_state['status'] = 'error'
            bot_state['error'] = 'Desktop automation not installed or unsupported in cloud container. Please use Web WhatsApp Queue.'
            bot_state['status_text'] = 'Desktop automation unavailable. Please use Web Queue.'
            bot_state['running'] = False
        return

    db = SessionLocal()
    try:
        raw_items = []
        if custom_items and isinstance(custom_items, list) and len(custom_items) > 0:
            for itm in custom_items:
                raw_items.append({
                    'id': itm.get('id', 0),
                    'name': itm.get('name', 'Valued Customer'),
                    'phone': str(itm.get('phone', '')),
                    'message': itm.get('message', ''),
                    'image_path': itm.get('image_path') or global_image_path or '',
                    'whatsapp_status': itm.get('whatsapp_status', 'active')
                })
        else:
            today = datetime.date.today()
            today_md = today.strftime('%m-%d')
            today_str = today.strftime('%Y-%m-%d')

            if queue_type == 'birthday':
                tpl = db.query(MessageTemplate).filter_by(category='birthday', is_default=True).first()
                if not tpl:
                    tpl = db.query(MessageTemplate).filter_by(category='birthday').first()
                customers = db.query(Customer).filter(Customer.dob.like(f"%-{today_md}")).all()
                for c in customers:
                    msg = render_template_text(tpl.body if tpl else 'Happy Birthday {name}!', c)
                    raw_items.append({
                        'id': c.id,
                        'name': c.name,
                        'phone': c.phone,
                        'message': msg,
                        'image_path': global_image_path or '',
                        'whatsapp_status': c.whatsapp_status or 'active'
                    })
            elif queue_type == 'expiry':
                tpl = db.query(MessageTemplate).filter_by(category='expiry', is_default=True).first()
                if not tpl:
                    tpl = db.query(MessageTemplate).filter_by(category='expiry').first()
                ahead = (today + datetime.timedelta(days=180)).strftime('%Y-%m-%d')
                customers = db.query(Customer).filter(and_(Customer.expiry_date >= today_str, Customer.expiry_date <= ahead)).all()
                for c in customers:
                    msg = render_template_text(tpl.body if tpl else 'Passport Renewal Reminder for {name}.', c)
                    raw_items.append({
                        'id': c.id,
                        'name': c.name,
                        'phone': c.phone,
                        'message': msg,
                        'image_path': global_image_path or '',
                        'whatsapp_status': c.whatsapp_status or 'active'
                    })

        with bot_lock:
            bot_state['total'] = len(raw_items)
            bot_state['current'] = 0
            bot_state['sent_count'] = 0
            bot_state['skipped_count'] = 0
            bot_state['skipped_items'] = []
            bot_state['status'] = 'running'
            bot_state['status_text'] = f"Found {len(raw_items)} contact(s). Starting automated dispatch..."

        if not raw_items:
            with bot_lock:
                bot_state['running'] = False
                bot_state['status'] = 'finished'
                bot_state['status_text'] = 'No contacts found to send.'
            return

        for idx, item in enumerate(raw_items, 1):
            if bot_state.get('stop_requested'):
                with bot_lock:
                    bot_state['status'] = 'stopped'
                    bot_state['status_text'] = 'Queue stopped by user.'
                break

            # Anti-ban Cool-down: 25-second break every 25 messages
            if idx > 1 and (idx - 1) % 25 == 0:
                with bot_lock:
                    bot_state['status'] = 'cooling_down'
                for remaining_cool in range(25, 0, -1):
                    if bot_state.get('stop_requested'):
                        break
                    with bot_lock:
                        bot_state['status_text'] = f"🛡️ Safety Pause: Cooling down for {remaining_cool}s (Anti-Ban protection after 25 messages)..."
                    time.sleep(1)
                if bot_state.get('stop_requested'):
                    with bot_lock:
                        bot_state['status'] = 'stopped'
                        bot_state['status_text'] = 'Queue stopped by user.'
                    break

            with bot_lock:
                bot_state['current'] = idx
                bot_state['current_id'] = item['id']
                bot_state['current_name'] = item['name']
                bot_state['current_phone'] = item['phone']
                bot_state['current_message'] = item['message']
                bot_state['skip_requested'] = False

            # Check 1: Already marked as no_whatsapp in DB?
            if item.get('whatsapp_status') == 'no_whatsapp':
                with bot_lock:
                    bot_state['skipped_count'] += 1
                    bot_state['skipped_items'].append({
                        'id': item['id'],
                        'name': item['name'],
                        'phone': item['phone'],
                        'reason': 'Marked as No WhatsApp in CRM'
                    })
                    bot_state['status_text'] = f"Skipping {item['name']}: Marked as No WhatsApp"
                log_entry = SentLog(
                    customer_id=item['id'],
                    customer_name=item['name'],
                    phone=item['phone'],
                    message_type=queue_type,
                    message_content=item['message'],
                    status='skipped_no_whatsapp'
                )
                db.add(log_entry)
                db.commit()
                time.sleep(0.3)
                continue

            # Check 2: Pre-validate phone format (10 digits, starts with 6-9, not landline)
            is_valid, clean_p, reason = validate_mobile_for_queue(item['phone'])
            if not is_valid:
                with bot_lock:
                    bot_state['skipped_count'] += 1
                    bot_state['skipped_items'].append({
                        'id': item['id'],
                        'name': item['name'],
                        'phone': item['phone'],
                        'reason': reason
                    })
                    bot_state['status_text'] = f"Skipping {item['name']}: {reason}"
                log_entry = SentLog(
                    customer_id=item['id'],
                    customer_name=item['name'],
                    phone=item['phone'],
                    message_type=queue_type,
                    message_content=item['message'],
                    status='skipped_invalid'
                )
                db.add(log_entry)
                db.commit()
                time.sleep(0.5)
                continue

            with bot_lock:
                bot_state['status'] = 'loading_whatsapp'
                bot_state['status_text'] = f"[{idx}/{len(raw_items)}] Opening WhatsApp for {item['name']}..."

            # Randomized Human Jitter Delay (delay ± random jitter)
            jitter_delay = max(4.0, float(delay_sec) + random.uniform(-0.5, 1.2))

            has_image = bool(item.get('image_path') and os.path.exists(item['image_path']))

            if has_image:
                # Open direct chat without prefilled text in URL; text will be pasted into image caption
                whatsapp_url = f"https://web.whatsapp.com/send?phone={clean_p}"
            else:
                encoded_msg = urllib.parse.quote(item['message'])
                whatsapp_url = f"https://web.whatsapp.com/send?phone={clean_p}&text={encoded_msg}"

            # Open browser tab directly into WhatsApp Web chat
            webbrowser.open(whatsapp_url)

            # Wait for WhatsApp Web to load with fast 0.2s poll loop so skip is instantaneous
            waited = 0.0
            skipped_early = False
            while waited < jitter_delay:
                if bot_state.get('stop_requested'):
                    break
                if bot_state.get('skip_requested'):
                    skipped_early = True
                    break
                time.sleep(0.2)
                waited += 0.2
                with bot_lock:
                    remaining = max(0, round(jitter_delay - waited, 1))
                    bot_state['status_text'] = f"Loading chat for {item['name']} ({remaining}s remaining)..."

            if bot_state.get('stop_requested'):
                with bot_lock:
                    bot_state['status'] = 'stopped'
                    bot_state['status_text'] = 'Queue stopped by user.'
                break

            if skipped_early:
                # User flagged that this contact is not on WhatsApp
                with bot_lock:
                    bot_state['skipped_count'] += 1
                    bot_state['skipped_items'].append({
                        'id': item['id'],
                        'name': item['name'],
                        'phone': item['phone'],
                        'reason': 'Skipped by user (No WhatsApp)'
                    })
                    bot_state['status'] = 'closing_tab'
                    bot_state['status_text'] = f"Marked {item['name']} as No WhatsApp. Closing tab..."

                # Update customer in DB
                cust_rec = db.query(Customer).filter_by(id=item['id']).first()
                if cust_rec:
                    cust_rec.whatsapp_status = 'no_whatsapp'
                log_entry = SentLog(
                    customer_id=item['id'],
                    customer_name=item['name'],
                    phone=item['phone'],
                    message_type=queue_type,
                    message_content=item['message'],
                    status='failed_no_whatsapp'
                )
                db.add(log_entry)
                db.commit()

                try:
                    pyautogui.hotkey('ctrl', 'w')
                except Exception:
                    pass
                time.sleep(0.5)
                continue

            # Send message or image
            with bot_lock:
                bot_state['status'] = 'sending'
                bot_state['status_text'] = f"Sending to {item['name']}..."

            try:
                if has_image:
                    # 1. Copy image to Windows clipboard via PowerShell
                    abs_img = os.path.abspath(item['image_path'])
                    subprocess.run(
                        ['powershell', '-NoProfile', '-Command', f'Set-Clipboard -Path "{abs_img}"'],
                        check=False,
                        creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
                    )
                    time.sleep(0.6)
                    # 2. Paste image into WhatsApp chat
                    pyautogui.hotkey('ctrl', 'v')
                    time.sleep(1.5) # Wait for image preview modal and caption field focus
                    # 3. If there is a caption message, copy to clipboard & paste
                    if item.get('message'):
                        pyperclip.copy(item['message'])
                        time.sleep(0.3)
                        pyautogui.hotkey('ctrl', 'v')
                        time.sleep(0.5)
                    # 4. Press Enter to send image with caption
                    pyautogui.press('enter')
                else:
                    pyautogui.press('enter')

                # If renewal alert, advance lead pipeline status
                if queue_type == 'expiry':
                    cust_rec = db.query(Customer).filter_by(id=item['id']).first()
                    if cust_rec and (not cust_rec.service_status or cust_rec.service_status == 'New'):
                        cust_rec.service_status = 'Alert Sent'

                # Log send to database
                log_entry = SentLog(
                    customer_id=item['id'],
                    customer_name=item['name'],
                    phone=item['phone'],
                    message_type=queue_type,
                    message_content=item['message'],
                    status='sent'
                )
                db.add(log_entry)
                db.commit()

                with bot_lock:
                    bot_state['sent_count'] += 1
                    bot_state['status_text'] = f"Message sent to {item['name']}! Closing tab..."
            except Exception as e:
                print(f"[Bot] Send error: {e}")

            time.sleep(random.uniform(0.8, 1.3)) # Humanized post-send pause

            if bot_state.get('stop_requested'):
                break

            # Close WhatsApp tab cleanly
            with bot_lock:
                bot_state['status'] = 'closing_tab'
            try:
                pyautogui.hotkey('ctrl', 'w')
            except Exception:
                pass

            # Safe human jitter pause before next contact (1.0s to 1.8s)
            with bot_lock:
                bot_state['status'] = 'waiting_next'
                bot_state['status_text'] = f"Tab closed. Preparing next contact..."

            time.sleep(random.uniform(1.0, 1.8))

        with bot_lock:
            if not bot_state.get('stop_requested') and bot_state.get('status') != 'stopped':
                bot_state['status'] = 'finished'
                sent = bot_state['sent_count']
                skp = bot_state['skipped_count']
                bot_state['status_text'] = f"Done! Sent: {sent}, Skipped/No WhatsApp: {skp} (Total: {len(raw_items)})."
            bot_state['running'] = False
            bot_state['stop_requested'] = False
            bot_state['skip_requested'] = False
    except Exception as e:
        with bot_lock:
            if 'FailSafeException' in type(e).__name__:
                bot_state['status'] = 'stopped'
                bot_state['error'] = 'Safety stop: Mouse moved to screen corner.'
                bot_state['status_text'] = 'Emergency Stop Activated: Mouse moved to corner.'
            else:
                bot_state['status'] = 'error'
                bot_state['error'] = str(e)
                bot_state['status_text'] = f"Error: {e}"
            bot_state['running'] = False
    finally:
        db.close()

@app.route('/api/bot/start', methods=['POST'])
@login_required
def start_bot():
    global bot_state
    with bot_lock:
        if bot_state['running']:
            return jsonify({'status': 'already_running', 'message': 'Bot is already running.'})
        data = request.json or {}
        qtype = data.get('type', 'birthday')
        delay = float(data.get('delay', 5.0))
        custom_items = data.get('items', None)
        image_path = (data.get('image_path') or '').strip()

        bot_state['running'] = True
        bot_state['stop_requested'] = False
        bot_state['skip_requested'] = False
        bot_state['type'] = qtype
        bot_state['status'] = 'starting'
        bot_state['status_text'] = 'Initializing automated queue...'
        bot_state['sent_count'] = 0
        bot_state['skipped_count'] = 0
        bot_state['skipped_items'] = []
        bot_state['total'] = len(custom_items) if custom_items else 0
        bot_state['current'] = 0
        bot_state['error'] = ''

    t = threading.Thread(target=bot_worker_thread, args=(qtype, delay, custom_items, image_path), daemon=True)
    t.start()
    return jsonify({'status': 'started'})

@app.route('/api/bot/skip', methods=['POST'])
@login_required
def skip_bot():
    global bot_state
    with bot_lock:
        if not bot_state['running']:
            return jsonify({'status': 'not_running', 'message': 'Bot is not running.'})
        bot_state['skip_requested'] = True
        bot_state['status_text'] = f"Skipping {bot_state.get('current_name', 'current contact')} (No WhatsApp)..."
    return jsonify({'status': 'success', 'message': 'Skip requested'})

@app.route('/api/bot/status', methods=['GET'])
@login_required
def get_bot_status():
    with bot_lock:
        return jsonify(bot_state)

@app.route('/api/bot/stop', methods=['POST'])
@login_required
def stop_bot():
    global bot_state
    with bot_lock:
        bot_state['stop_requested'] = True
        bot_state['status'] = 'stopped'
        bot_state['status_text'] = 'Stopping bot...'
        bot_state['running'] = False
    return jsonify({'status': 'stopping'})

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print(f"🚀 PassWish Server is starting at http://127.0.0.1:{port}")
    app.run(host='0.0.0.0', port=port, debug=False, threaded=True)
